# -*- coding: utf-8 -*-
"""Convertit le rapport d'audit Markdown en DOCX (python-docx)."""
import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT

SRC = '/home/user/Chansons/Audit_rapport_quantitatif_SACO_Phase2_2026-10-20.md'
OUT = '/home/user/Chansons/Audit_rapport_quantitatif_SACO_Phase2_2026-10-20.docx'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin = sec.bottom_margin = Cm(2.0)

st = doc.styles['Normal']
st.font.name = 'Calibri'; st.font.size = Pt(10)
st.paragraph_format.space_after = Pt(4)
st.paragraph_format.line_spacing = 1.08

for name, size, color in (('Heading 1', 17, (0x1F, 0x3A, 0x5F)),
                          ('Heading 2', 13, (0x1F, 0x3A, 0x5F)),
                          ('Heading 3', 11.5, (0x2E, 0x4E, 0x73))):
    s = doc.styles[name]
    s.font.name = 'Calibri'; s.font.size = Pt(size); s.font.bold = True
    s.font.color.rgb = RGBColor(*color)
    s.paragraph_format.space_before = Pt(14); s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.keep_with_next = True

INLINE = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)')

def add_runs(par, text):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith('`') and tok.endswith('`'):
            r = par.add_run(tok[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(8.5)
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.italic = True
        else:
            par.add_run(tok)

def add_table(rows):
    header = [c.strip() for c in rows[0].strip('|').split('|')]
    body = []
    for line in rows[2:]:
        body.append([c.strip() for c in line.strip('|').split('|')])
    ncol = len(header)
    t = doc.add_table(rows=1, cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(header):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]
        r = p.add_run(h.replace('**', '')); r.bold = True; r.font.size = Pt(8.5)
        p.paragraph_format.space_after = Pt(1)
    for row in body:
        cells = t.add_row().cells
        for i in range(ncol):
            val = row[i] if i < len(row) else ''
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            add_runs(p, val)
            for r in p.runs:
                if r.font.size is None:
                    r.font.size = Pt(8)
    doc.add_paragraph()

lines = open(SRC, encoding='utf-8').read().split('\n')
i = 0
first_h1 = True
while i < len(lines):
    line = lines[i].rstrip()
    if not line.strip():
        i += 1; continue
    if line.startswith('|'):
        block = []
        while i < len(lines) and lines[i].lstrip().startswith('|'):
            block.append(lines[i].strip()); i += 1
        if len(block) >= 2:
            add_table(block)
        continue
    if line.startswith('### '):
        doc.add_heading(line[4:].strip(), level=3)
    elif line.startswith('## '):
        doc.add_heading(line[3:].strip(), level=2)
    elif line.startswith('# '):
        txt = line[2:].strip()
        if not first_h1:
            doc.add_page_break()
        first_h1 = False
        doc.add_heading(txt, level=1)
    elif line.startswith('---'):
        pass
    elif line.startswith('> '):
        txt = line[2:].strip()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_before = Pt(8)
        add_runs(p, txt)
        for r in p.runs:
            r.bold = True
            r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    else:
        p = doc.add_paragraph()
        if line.startswith('*') and line.endswith('*') and not line.startswith('**'):
            p.paragraph_format.left_indent = Cm(0.3)
        add_runs(p, line)
    i += 1

doc.save(OUT)
print('écrit', OUT)
