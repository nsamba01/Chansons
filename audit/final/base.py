# -*- coding: utf-8 -*-
"""Assemblage du rapport final sur la maquette Word.

La maquette sert de base : styles, en-têtes, pieds de page, largeurs de page et
gabarits de tableaux sont repris tels quels. Le corps est reconstruit : les
douze blocs « À rédiger » sont remplacés par le texte, les figures et les
tableaux définitifs.
"""
import copy, json, os, sys
from docx import Document
from docx.shared import Pt, Cm, Emu, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = '/home/user/Chansons'
TPL = os.path.join(ROOT, 'Rapport_resultats_quantitatifs_Phase2_SACO_Mpox_maquette_2026-10-19 (1).docx')
FIG = os.path.join(HERE, 'fig')
X = json.load(open('/home/user/dump.json'))

NAVY = '0B1F3A'
BODY = '10233F'
CYAN = '22B8CF'
GREY = '6B7280'
RULE = 'D9E1EC'


# ------------------------------------------------------------------ bas niveau
def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:shd')):
        tcPr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def cell_margins(cell, l=70, r=70, t=40, b=40):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for k, v in (('top', t), ('left', l), ('bottom', b), ('right', r)):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tcPr.append(mar)


def borders(table, top=NAVY, bottom=NAVY, inside='D9E1EC', sz=8, isz=4):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for tag, color, s in (('top', top, sz), ('left', 'FFFFFF', 0), ('bottom', bottom, sz),
                          ('right', 'FFFFFF', 0), ('insideH', inside, isz), ('insideV', inside, isz)):
        e = OxmlElement(f'w:{tag}')
        e.set(qn('w:val'), 'single' if s else 'none')
        e.set(qn('w:sz'), str(s)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        b.append(e)
    tblPr.append(b)


def keep_together(table):
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        e = OxmlElement('w:cantSplit'); trPr.append(e)


def header_repeat(table):
    trPr = table.rows[0]._tr.get_or_add_trPr()
    e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); trPr.append(e)


def set_width(table, cm):
    table.autofit = False
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn('w:tblW')):
        tblPr.remove(old)
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), str(int(cm * 567))); w.set(qn('w:type'), 'dxa')
    tblPr.append(w)
    for row in table.rows:
        for c in row.cells:
            c.width = Cm(cm / len(row.cells))


def run(p, text, size=10, bold=False, color=BODY, italic=False, font='Calibri'):
    r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    r.font.color.rgb = RGBColor.from_string(color); r.font.name = font
    return r


def para(doc, text='', size=10, bold=False, color=BODY, italic=False, after=100,
         line=276, align='both', before=0, keep=False, style=None):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.space_after = Pt(after / 20.0); pf.space_before = Pt(before / 20.0)
    pf.line_spacing = Pt(line / 20.0)
    if align == 'both':
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == 'center':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'left':
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if keep:
        pf.keep_with_next = True
    if text:
        # gras encadré par ** dans le texte
        for i, chunk in enumerate(text.split('**')):
            if chunk:
                run(p, chunk, size=size, bold=bold or (i % 2 == 1), color=color, italic=italic)
    return p


def md_runs(p, text, size=10, color=BODY):
    """Texte avec **gras** et *italique*."""
    import re
    for tok in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text):
        if not tok:
            continue
        if tok.startswith('**'):
            run(p, tok[2:-2], size=size, bold=True, color=color)
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            run(p, tok[1:-1], size=size, italic=True, color=color)
        else:
            run(p, tok, size=size, color=color)


def _style(p, style_id):
    pPr = p._p.get_or_add_pPr()
    for old in pPr.findall(qn('w:pStyle')):
        pPr.remove(old)
    e = OxmlElement('w:pStyle'); e.set(qn('w:val'), style_id)
    pPr.insert(0, e)
    return p


def h1(doc, number, title, band=None):
    """Titre de chapitre : bandeau pleine largeur + Heading 1 invisible (pour la TDM)."""
    if band and os.path.exists(band):
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.paragraph_format.space_after = Pt(2)
        pic.paragraph_format.space_before = Pt(0)
        pic.add_run().add_picture(band, width=Cm(17.0))
    p = _style(doc.add_paragraph(), 'Heading1')
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    run(p, f'{number}. {title}' if number else title, size=1, color='FFFFFF')
    return p


def h2(doc, text, number=None):
    p = _style(doc.add_paragraph(), 'Heading2')
    p.paragraph_format.space_before = Pt(13); p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    run(p, (f'{number} {text}' if number else text), size=13, bold=True, color=NAVY)
    return p


def h3(doc, text):
    p = _style(doc.add_paragraph(), 'Heading3')
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run(p, text, size=10.5, bold=True, color='2E74B5')
    return p


def bullets(doc, items, size=10, numbered=False):
    for it in items:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.55); pf.first_line_indent = Cm(-0.35)
        pf.space_after = Pt(1.5); pf.line_spacing = Pt(13.2)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run(p, '•  ' if not numbered else '', size=size, color=CYAN, bold=True)
        md_runs(p, it, size=size)


def caption(doc, text, before=6, after=2):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.keep_with_next = True
    run(p, text, size=8.5, bold=True, color=NAVY)
    return p


def note(doc, text, size=7.3):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(1); pf.space_after = Pt(9); pf.line_spacing = Pt(9.3)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    md_runs(p, text, size=size, color=GREY)


def table(doc, caption_txt, head, rows, widths=None, note_txt=None, size=7.6, head_size=6.6,
          first_col_left=True):
    caption(doc, caption_txt)
    t = doc.add_table(rows=1, cols=len(head))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0]
    for i, h in enumerate(head):
        c = hdr.cells[i]
        shade(c, NAVY); cell_margins(c)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(8.2)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
        run(p, str(h), size=head_size, bold=True, color='FFFFFF')
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            c = cells[i]; cell_margins(c)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(size + 1.1)
            p.alignment = (WD_ALIGN_PARAGRAPH.LEFT if (i == 0 and first_col_left)
                           else WD_ALIGN_PARAGRAPH.CENTER)
            txt = str(v)
            bold = txt.startswith('*') and txt.endswith('*')
            if bold:
                txt = txt.strip('*')
            run(p, txt, size=size, bold=bold, color=BODY)
        if len(r) and str(r[-1]).strip() in ('Sous-total', 'Total'):
            pass
    total_w = 17.0
    if widths:
        ssum = sum(widths)
        if ssum > total_w:
            widths = [w * total_w / ssum for w in widths]
        t.autofit = False
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    set_width(t, total_w) if not widths else None
    borders(t); keep_together(t); header_repeat(t)
    if note_txt:
        note(doc, note_txt)


def figure(doc, path, caption_txt, width=16.6, note_txt=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(width))
    caption(doc, caption_txt)
    if note_txt:
        note(doc, note_txt)


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def hr(doc, width=17.0, color=CYAN):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(8)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '8'); b.set(qn('w:space'), '1'); b.set(qn('w:color'), color)
    pbdr.append(b); pPr.append(pbdr)


def callout(doc, title, text, color=NAVY, fill='F2F6FA'):
    """Encadré « ce qu'il faut retenir »."""
    t = doc.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    shade(c, fill); cell_margins(c, l=140, r=140, t=90, b=90)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run(p, title, size=8.4, bold=True, color=color)
    for chunk in text if isinstance(text, list) else [text]:
        p2 = c.add_paragraph()
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.line_spacing = Pt(11.4)
        p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        md_runs(p2, chunk, size=8.4)
    set_width(t, 17.0)
    # bordure gauche colorée uniquement
    tblPr = t._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for tag, s, col in (('top', 0, 'FFFFFF'), ('bottom', 0, 'FFFFFF'), ('left', 24, color),
                        ('right', 0, 'FFFFFF'), ('insideH', 0, 'FFFFFF'), ('insideV', 0, 'FFFFFF')):
        e = OxmlElement(f'w:{tag}')
        e.set(qn('w:val'), 'single' if s else 'none'); e.set(qn('w:sz'), str(s))
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), col)
        b.append(e)
    tblPr.append(b)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ------------------------------------------------------------------ numérotation
def _french(v):
    if isinstance(v, (int, float)):
        s = f'{v:g}'
        return s.replace('.', ',')
    return str(v)


def sheet_rows(name, cols, start=1, end=None, skip_blank=None):
    ws = X[name]['rows']
    out = []
    for r in ws[start:end]:
        vals = [r['cells'].get(c, '') for c in cols]
        if skip_blank and all(str(v).strip() in ('', '0') for v in vals):
            continue
        out.append([_french(v) for v in vals])
    return out


# ------------------------------------------------------------------ document
def load_template():
    doc = Document(TPL)
    body = doc.element.body
    kids = list(body)
    # conserver le secteur final (marges 2 cm, en-tête, pied de page)
    final_sect = kids[-1] if kids[-1].tag == qn('w:sectPr') else None
    if final_sect is not None:
        body.remove(final_sect)
    # garder : couverture (p0) + paragraphe porteur du secteur 1 (p1)
    for e in kids[2:]:
        if e is not final_sect:
            body.remove(e)
    if final_sect is not None:
        body.append(final_sect)
    return doc


def set_update_fields(doc):
    settings = doc.settings.element
    for old in settings.findall(qn('w:updateFields')):
        settings.remove(old)
    e = OxmlElement('w:updateFields'); e.set(qn('w:val'), 'true')
    settings.append(e)
